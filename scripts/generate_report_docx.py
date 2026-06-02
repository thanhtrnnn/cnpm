"""Generate DOCX for BÁO CÁO (Phần I–V).

Sources: docs/tabs/report/phan-*.md
Output:  output/report.docx

Font:     Times New Roman throughout
Headings: TNR, bold, #1F4E79 (XÁC ĐỊNH YÊU CẦU style)
  # (PHẦN)  → 16pt
  ## (N.)   → 14pt
  ### (N.N.)→ 13pt
"""
import os
import re
from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import nsdecls
from docx.oxml import parse_xml

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
REPO_DIR   = os.path.join(SCRIPT_DIR, '..')
REPORT_DIR = os.path.join(REPO_DIR, 'docs', 'tabs', 'report')
OUTPUT_DIR = os.path.join(REPO_DIR, 'output')
OUTPUT_FILE = os.path.join(OUTPUT_DIR, 'report.docx')

# UC overview diagrams — identified by visual inspection of screenshots
DIAGRAM_MAP = {
    # UC overviews
    'account_uc_overview': os.path.join(REPO_DIR, 'docs', 'tabs', 'exports', 'account',  'screenshots', 'image_01.png'),
    'booking_uc_overview':  os.path.join(REPO_DIR, 'docs', 'tabs', 'exports', 'booking',  'screenshots', 'image_01.png'),
    'services_uc_overview': os.path.join(REPO_DIR, 'docs', 'tabs', 'exports', 'services', 'screenshots', 'image_01.png'),
    'core_uc_overview':     os.path.join(REPO_DIR, 'docs', 'tabs', 'exports', 'core',     'screenshots', 'image_01.png'),
    'hr_uc_overview':       os.path.join(REPO_DIR, 'docs', 'tabs', 'exports', 'report',   'screenshots', 'image_01.png'),
    # Entity diagrams
    'account_entity':  os.path.join(REPO_DIR, 'docs', 'tabs', 'exports', 'account',  'screenshots', 'image_07.png'),
    'booking_entity':  os.path.join(REPO_DIR, 'docs', 'tabs', 'exports', 'booking',  'screenshots', 'image_06.png'),
    'services_entity': os.path.join(REPO_DIR, 'docs', 'tabs', 'exports', 'services', 'screenshots', 'image_06.png'),
    'core_entity':     os.path.join(REPO_DIR, 'docs', 'tabs', 'exports', 'core',     'screenshots', 'image_06.png'),
    'hr_entity':       os.path.join(REPO_DIR, 'docs', 'tabs', 'exports', 'report',   'screenshots', 'image_06.png'),
    # BCE class diagrams
    'account_bce':  os.path.join(REPO_DIR, 'docs', 'tabs', 'exports', 'account',  'screenshots', 'image_08.png'),
    'booking_bce':  os.path.join(REPO_DIR, 'docs', 'tabs', 'exports', 'booking',  'screenshots', 'image_07.png'),
    'services_bce': os.path.join(REPO_DIR, 'docs', 'tabs', 'exports', 'services', 'screenshots', 'image_07.png'),
    'core_bce':     os.path.join(REPO_DIR, 'docs', 'tabs', 'exports', 'core',     'screenshots', 'image_07.png'),
    'hr_bce':       os.path.join(REPO_DIR, 'docs', 'tabs', 'exports', 'report',   'screenshots', 'image_07.png'),
    # Sequence diagrams
    'account_seq':  os.path.join(REPO_DIR, 'docs', 'tabs', 'exports', 'account',  'screenshots', 'image_09.png'),
    'booking_seq':  os.path.join(REPO_DIR, 'docs', 'tabs', 'exports', 'booking',  'screenshots', 'image_11.png'),
    'services_seq': os.path.join(REPO_DIR, 'docs', 'tabs', 'exports', 'services', 'screenshots', 'image_11.png'),
    'core_seq':     os.path.join(REPO_DIR, 'docs', 'tabs', 'exports', 'core',     'screenshots', 'image_12.png'),
    'hr_seq':       os.path.join(REPO_DIR, 'docs', 'tabs', 'exports', 'report',   'screenshots', 'image_08.png'),
}

# Markdown files to process in order
MD_FILES = [
    os.path.join(REPORT_DIR, 'phan-i-ii.md'),
    os.path.join(REPORT_DIR, 'phan-iii.md'),
    os.path.join(REPORT_DIR, 'phan-iv.md'),
    os.path.join(REPORT_DIR, 'phan-v.md'),
]

TNR = 'Times New Roman'
HEADING_COLOR = RGBColor(0x1F, 0x4E, 0x79)
HEADING_SIZES = {1: Pt(16), 2: Pt(14), 3: Pt(13), 4: Pt(12)}
BODY_SIZE = Pt(12)
TABLE_SIZE = Pt(11)


# ─── inline formatting ────────────────────────────────────────────────────────

def _parse_inline(text):
    """Return list of (text, bold, italic, code) segments."""
    segs = []
    i = 0
    buf = ''
    while i < len(text):
        # bold **
        if text[i:i+2] == '**':
            if buf:
                segs.append((buf, False, False, False))
                buf = ''
            j = text.find('**', i + 2)
            if j != -1:
                segs.append((text[i+2:j], True, False, False))
                i = j + 2
                continue
        # italic * (single, not **)
        elif text[i] == '*' and text[i:i+2] != '**':
            if buf:
                segs.append((buf, False, False, False))
                buf = ''
            j = text.find('*', i + 1)
            if j != -1 and text[j:j+2] != '**':
                segs.append((text[i+1:j], False, True, False))
                i = j + 1
                continue
        # code `
        elif text[i] == '`':
            if buf:
                segs.append((buf, False, False, False))
                buf = ''
            j = text.find('`', i + 1)
            if j != -1:
                segs.append((text[i+1:j], False, False, True))
                i = j + 1
                continue
        buf += text[i]
        i += 1
    if buf:
        segs.append((buf, False, False, False))
    return segs


def _apply_runs(p, text, size=BODY_SIZE):
    for seg, bold, italic, code in _parse_inline(text):
        run = p.add_run(seg)
        run.font.name = TNR
        run.font.size = size
        run.bold = bold
        run.italic = italic
        if code:
            run.font.name = 'Courier New'
            run.font.size = Pt(10)


# ─── block helpers ────────────────────────────────────────────────────────────

def add_heading(doc, text, level):
    p = doc.add_heading(level=level)
    size = HEADING_SIZES.get(level, Pt(12))
    for seg, bold, italic, code in _parse_inline(text):
        run = p.add_run(seg)
        run.font.name = TNR
        run.font.size = size
        run.font.bold = True
        run.font.color.rgb = HEADING_COLOR
        if italic:
            run.italic = True
        if code:
            run.font.name = 'Courier New'
    return p


def add_para(doc, text, size=BODY_SIZE):
    p = doc.add_paragraph()
    p.style = doc.styles['Normal']
    _apply_runs(p, text, size)
    return p


def _set_shading(cell, color):
    shading = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color}"/>')
    cell._tc.get_or_add_tcPr().append(shading)


def add_table(doc, rows):
    if not rows:
        return
    num_cols = max(len(r) for r in rows)
    tbl = doc.add_table(rows=len(rows), cols=num_cols)
    tbl.style = 'Table Grid'
    tbl.alignment = WD_TABLE_ALIGNMENT.LEFT
    for ri, row_data in enumerate(rows):
        for ci in range(num_cols):
            cell = tbl.cell(ri, ci)
            cell.text = ''
            p = cell.paragraphs[0]
            p.style = doc.styles['Normal']
            cell_text = row_data[ci] if ci < len(row_data) else ''
            _apply_runs(p, cell_text, TABLE_SIZE)
            if ri == 0:
                for run in p.runs:
                    run.bold = True
                _set_shading(cell, 'D9E2F3')
    return tbl


def _setup_bullets(doc):
    numbering_xml = doc.part.numbering_part._element
    W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    abst = OxmlElement('w:abstractNum')
    abst.set(W + 'abstractNumId', '1')
    ml = OxmlElement('w:multiLevelType')
    ml.set(W + 'val', 'hybridMultilevel')
    abst.append(ml)
    for idx, (left, hang) in enumerate([(360, 360), (720, 360), (1080, 360)]):
        lvl = OxmlElement('w:lvl')
        lvl.set(W + 'ilvl', str(idx))
        st = OxmlElement('w:start'); st.set(W + 'val', '1'); lvl.append(st)
        fmt = OxmlElement('w:numFmt'); fmt.set(W + 'val', 'bullet'); lvl.append(fmt)
        sym = '•' if idx == 0 else ('◦' if idx == 1 else '▪')
        lt = OxmlElement('w:lvlText'); lt.set(W + 'val', sym); lvl.append(lt)
        pPr = OxmlElement('w:pPr')
        ind = OxmlElement('w:ind')
        ind.set(W + 'left', str(left)); ind.set(W + 'hanging', str(hang))
        pPr.append(ind); lvl.append(pPr)
        abst.append(lvl)
    numbering_xml.append(abst)
    num = OxmlElement('w:num'); num.set(W + 'numId', '1')
    ref = OxmlElement('w:abstractNumId'); ref.set(W + 'val', '1')
    num.append(ref); numbering_xml.append(num)


def add_bullet(doc, text, level=0):
    p = add_para(doc, text)
    W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
    pPr = p._element.get_or_add_pPr()
    numPr = OxmlElement('w:numPr')
    ilvl = OxmlElement('w:ilvl'); ilvl.set(W + 'val', str(level)); numPr.append(ilvl)
    numId = OxmlElement('w:numId'); numId.set(W + 'val', '1'); numPr.append(numId)
    pPr.append(numPr)
    return p


# ─── markdown processor ───────────────────────────────────────────────────────

def process_md(doc, md_path):
    with open(md_path, encoding='utf-8') as f:
        lines = f.read().split('\n')

    i = 0
    in_table = False
    current_table = []

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # PLACEHOLDER comment → embed diagram image
        if stripped.startswith('<!--') and 'PLACEHOLDER' in stripped:
            m_ph = re.search(r'PLACEHOLDER:\s*(\S+)', stripped)
            if m_ph:
                key = m_ph.group(1).rstrip('-->')  .strip()
                img_path = DIAGRAM_MAP.get(key)
                if img_path and os.path.exists(img_path):
                    p = doc.add_paragraph()
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    run = p.add_run()
                    run.add_picture(img_path, width=Inches(5.5))
                else:
                    add_para(doc, f'[Biểu đồ: {key}]')
            i += 1
            continue

        # Other HTML comments — skip entirely
        if stripped.startswith('<!--'):
            i += 1
            continue

        # Horizontal rule --- (UC separators)
        if re.match(r'^-{3,}$', stripped):
            i += 1
            continue

        # Empty line
        if not stripped:
            if in_table and current_table:
                add_table(doc, current_table)
                current_table = []
                in_table = False
            i += 1
            continue

        # Table separator line |---|---|
        if re.match(r'^\|[\s\-:|]+\|$', stripped):
            i += 1
            continue

        # Table row
        if stripped.startswith('|') and stripped.endswith('|'):
            if not in_table:
                in_table = True
            cells = [c.strip() for c in stripped.split('|')[1:-1]]
            current_table.append(cells)
            i += 1
            continue

        # Flush pending table before non-table content
        if in_table and current_table:
            add_table(doc, current_table)
            current_table = []
            in_table = False

        # Heading  # / ## / ### / ####
        m = re.match(r'^(#{1,4})\s+(.+)$', line)
        if m:
            level = len(m.group(1))
            add_heading(doc, m.group(2).strip(), level)
            i += 1
            continue

        # Bullet  - text  (any indent level)
        bm = re.match(r'^(\s*)[-*]\s+(.+)$', line)
        if bm:
            level = min(len(bm.group(1)) // 2, 2)
            add_bullet(doc, bm.group(2).strip(), level)
            i += 1
            continue

        # Regular paragraph (includes **bold** standalone, italic labels, flow text)
        add_para(doc, stripped)
        i += 1

    # Flush any remaining table
    if in_table and current_table:
        add_table(doc, current_table)


# ─── main ─────────────────────────────────────────────────────────────────────

def main():
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    doc = Document()

    _setup_bullets(doc)

    # Global styles
    normal = doc.styles['Normal']
    normal.font.name = TNR
    normal.font.size = BODY_SIZE

    for lvl in range(1, 5):
        hs = doc.styles[f'Heading {lvl}']
        hs.font.name = TNR
        hs.font.size = HEADING_SIZES[lvl]
        hs.font.bold = True
        hs.font.color.rgb = HEADING_COLOR

    for md_path in MD_FILES:
        if not os.path.exists(md_path):
            print(f'SKIP (not found): {md_path}')
            continue
        print(f'Processing: {os.path.basename(md_path)}')
        process_md(doc, md_path)

    doc.save(OUTPUT_FILE)
    print(f'Saved: {OUTPUT_FILE}  ({os.path.getsize(OUTPUT_FILE):,} bytes)')


if __name__ == '__main__':
    main()
